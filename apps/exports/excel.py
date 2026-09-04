"""
Excel eksporti.

TZ talablari:
  * yaratilgan ID kodlar adminga **Excel fayl** ko'rinishida beriladi
    (№ | ID kod | Holati);
  * natijalar Excel/PDF ko'rinishida eksport qilinadi (SRS 9-bo'lim).

Fayllar `openpyxl` yordamida to'g'ridan-to'g'ri xotirada yig'iladi va
bayt massivi sifatida qaytariladi — diskka yozish shart emas.
"""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.accesscodes.models import AccessCode, CodeBatch
from apps.attempts.models import Attempt
from apps.attempts.services import ranked_attempts
from apps.exams.models import Exam, Question
from core import constants as C
from django.utils import timezone

# --------------------------------------------------------------------------
#  Umumiy uslublar
# --------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="0F2B46")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="0F2B46")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center")

THIN = Side(style="thin", color="BFC9D4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

STATUS_FILLS = {
    AccessCode.Status.UNUSED: PatternFill("solid", fgColor="E8F5E9"),
    AccessCode.Status.ACTIVATED: PatternFill("solid", fgColor="FFF8E1"),
    AccessCode.Status.USED: PatternFill("solid", fgColor="FFEBEE"),
    AccessCode.Status.REVOKED: PatternFill("solid", fgColor="ECEFF1"),
}


def _style_header(worksheet, row: int, columns: int) -> None:
    """Sarlavha qatorini bezaydi."""
    for column in range(1, columns + 1):
        cell = worksheet.cell(row=row, column=column)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def _auto_width(worksheet, widths: list[int]) -> None:
    """Ustun kengliklarini belgilaydi."""
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width


def _save(workbook: Workbook) -> bytes:
    """Kitobni bayt massiviga aylantiradi."""
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _timestamp() -> str:
    return timezone.localtime().strftime("%d.%m.%Y %H:%M")


# --------------------------------------------------------------------------
#  1. ID kodlar
# --------------------------------------------------------------------------


def codes_workbook(batch: CodeBatch) -> bytes:
    """
    ID kodlar ro'yxatini Excel ko'rinishida qaytaradi.

    Ustunlar: № | ID kod | Holati | Foydalanuvchi | Telegram ID | Ishlatilgan vaqt
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ID kodlar"

    sheet["A1"] = f"«{batch.exam.title}» — ID kodlar"
    sheet["A1"].font = TITLE_FONT
    sheet.merge_cells("A1:F1")
    sheet["A2"] = f"Test kodi: {batch.exam.code}   |   Yaratilgan: {_timestamp()}"
    sheet.merge_cells("A2:F2")

    headers = ["№", "ID kod", "Holati", "Ism-familiya", "Telegram ID", "Ishlatilgan vaqt"]
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=4, column=index, value=header)
    _style_header(sheet, 4, len(headers))

    codes = batch.codes.select_related("user").order_by("id")
    for row_index, code in enumerate(codes, start=1):
        row = row_index + 4
        used_at = (
            timezone.localtime(code.used_at).strftime("%d.%m.%Y %H:%M")
            if code.used_at
            else ""
        )
        values = [
            row_index,
            code.code,
            code.get_status_display(),
            code.full_name or "",
            code.telegram_id or "",
            used_at,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.border = BORDER
            cell.alignment = CENTER if column in {1, 2, 3, 5} else LEFT
        fill = STATUS_FILLS.get(code.status)
        if fill:
            sheet.cell(row=row, column=3).fill = fill

    _auto_width(sheet, [6, 16, 18, 28, 16, 20])
    sheet.freeze_panes = "A5"
    return _save(workbook)


# --------------------------------------------------------------------------
#  2. Natijalar
# --------------------------------------------------------------------------


def overall_results_workbook(exam: Exam) -> bytes:
    """
    E'lon uchun soddalashtirilgan Excel jadvali.

    Ustunlar: `№ · F.I.SH (yoki ID raqami) · Ball · Foiz · Daraja` va
    RASH testlarida fan ballari. Nechta savolni to'g'ri topgani, javoblar
    matritsasi va savollar statistikasi **bu faylda yo'q** — shu sababli
    uni testni yaratgan foydalanuvchiga ham berish mumkin. To'liq jadval
    `results_workbook()` da, u faqat asosiy adminlar uchun.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Natijalar"

    uses_rasch = exam.uses_rasch
    headers = ["№", "F.I.SH", "Ball", "Foiz", "Daraja"]
    widths = [6, 30, 10, 10, 16]
    if uses_rasch:
        headers += C.subject_names()
        widths += [14] * len(C.SUBJECT_SCORES)
    headers.append("Topshirgan vaqt")
    widths.append(18)

    sheet["A1"] = f"«{exam.title}» — natijalar"
    sheet["A1"].font = TITLE_FONT
    sheet.merge_cells(f"A1:{get_column_letter(len(headers))}1")
    sheet["A2"] = (
        f"Test kodi: {exam.code}   |   Turi: {exam.get_exam_type_display()}   |   "
        f"Eksport: {_timestamp()}"
    )
    sheet.merge_cells(f"A2:{get_column_letter(len(headers))}2")

    for index, header in enumerate(headers, start=1):
        sheet.cell(row=4, column=index, value=header)
    _style_header(sheet, 4, len(headers))

    for row_index, attempt in enumerate(ranked_attempts(exam), start=1):
        row = row_index + 4
        percent = (
            C.certificate_percent(attempt.ball, attempt.grade)
            if uses_rasch
            else round(attempt.percent, 2)
        )
        values = [
            attempt.rank or row_index,
            attempt.public_label,
            round(attempt.ball, 2) if attempt.ball is not None else "",
            percent,
            attempt.grade or "",
        ]
        if uses_rasch:
            values += C.subject_scores(attempt.ball, attempt.grade)
        values.append(
            timezone.localtime(attempt.submitted_at).strftime("%d.%m.%Y %H:%M")
            if attempt.submitted_at
            else ""
        )
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.border = BORDER
            cell.alignment = LEFT if column == 2 else CENTER

    _auto_width(sheet, widths)
    sheet.freeze_panes = "A5"
    return _save(workbook)


def results_workbook(exam: Exam) -> bytes:
    """Test natijalari va reytingini Excel ko'rinishida qaytaradi."""
    workbook = Workbook()

    # --- 1-varaq: reyting ---
    sheet = workbook.active
    sheet.title = "Reyting"

    sheet["A1"] = f"«{exam.title}» — natijalar"
    sheet["A1"].font = TITLE_FONT
    sheet.merge_cells("A1:J1")
    sheet["A2"] = (
        f"Test kodi: {exam.code}   |   Turi: {exam.get_exam_type_display()}   |   "
        f"Eksport: {_timestamp()}"
    )
    sheet.merge_cells("A2:J2")

    uses_rasch = exam.uses_rasch
    is_paid = exam.exam_type == Exam.Type.RASCH_PAID
    headers = ["O‘rin", "Ism-familiya", "Telefon", "Telegram ID"]
    # Pullik testda natijalar ID raqami bilan e'lon qilinadi — admin uchun
    # ikkala ustun ham kerak.
    if is_paid:
        headers.append("ID raqami")
    headers += ["To‘g‘ri", "Xato", "Bo‘sh", "Foiz"]
    if uses_rasch:
        # «Sert. foizi» — `ball * 100 / 65`. Asosiy fanlar shu foizga
        # proporsional, majburiy fan esa sertifikat olganlarning hammasiga
        # to'liq beriladi: 11 ball (`core.constants.subject_scores`).
        headers += ["theta", "Ball", "Daraja", "Sert. foizi"]
        headers += C.subject_names()
    headers += ["Topshirgan vaqt", "Sarflangan vaqt"]

    for index, header in enumerate(headers, start=1):
        sheet.cell(row=4, column=index, value=header)
    _style_header(sheet, 4, len(headers))

    attempts = ranked_attempts(exam)
    for row_index, attempt in enumerate(attempts, start=1):
        row = row_index + 4
        submitted = (
            timezone.localtime(attempt.submitted_at).strftime("%d.%m.%Y %H:%M")
            if attempt.submitted_at
            else ""
        )
        duration = f"{attempt.duration_seconds // 60} daq {attempt.duration_seconds % 60} s"
        values = [
            attempt.rank or row_index,
            attempt.full_name or attempt.user.display_name,
            attempt.phone or "",
            attempt.user.telegram_id,
        ]
        if is_paid:
            values.append(attempt.access_code_value or "—")
        values += [
            attempt.raw_score,
            attempt.wrong_count,
            attempt.empty_count,
            round(attempt.percent, 2),
        ]
        if uses_rasch:
            values += [
                round(attempt.theta, 4) if attempt.theta is not None else "",
                round(attempt.ball, 2) if attempt.ball is not None else "",
                attempt.grade or "",
                C.certificate_percent(attempt.ball, attempt.grade),
            ]
            values += C.subject_scores(attempt.ball, attempt.grade)
        values += [submitted, duration]

        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.border = BORDER
            cell.alignment = LEFT if column == 2 else CENTER

    widths = [8, 28, 16, 14]
    if is_paid:
        widths.append(14)
    widths += [9, 8, 8, 9]
    if uses_rasch:
        widths += [10, 10, 12, 11] + [13] * len(C.SUBJECT_SCORES)
    widths += [18, 16]
    _auto_width(sheet, widths)
    sheet.freeze_panes = "A5"

    # --- 2-varaq: javoblar matritsasi ---
    _add_answers_sheet(workbook, exam)

    # --- 3-varaq: savollar statistikasi ---
    _add_items_sheet(workbook, exam)

    return _save(workbook)


def _add_answers_sheet(workbook: Workbook, exam: Exam) -> None:
    """Har bir ishtirokchining har bir savolga bergan javobi."""
    sheet = workbook.create_sheet("Javoblar")
    questions = list(exam.questions.filter(is_active=True).order_by("order"))

    headers = ["Ism-familiya"] + [f"{q.order}" for q in questions] + ["Jami"]
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=index, value=header)
    _style_header(sheet, 1, len(headers))

    attempts = (
        Attempt.objects.filter(exam=exam, status=Attempt.Status.SUBMITTED)
        .select_related("user")
        .prefetch_related("answers")
        .ranked()
    )
    correct_fill = PatternFill("solid", fgColor="E8F5E9")
    wrong_fill = PatternFill("solid", fgColor="FFEBEE")

    for row_index, attempt in enumerate(attempts, start=2):
        sheet.cell(row=row_index, column=1, value=attempt.full_name or attempt.user.display_name)
        answers = {answer.question_id: answer for answer in attempt.answers.all()}
        for column_index, question in enumerate(questions, start=2):
            answer = answers.get(question.id)
            value = answer.display_value if answer else "—"
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.alignment = CENTER
            cell.border = BORDER
            if answer is not None and not answer.is_empty:
                cell.fill = correct_fill if answer.score > 0 else wrong_fill
        sheet.cell(
            row=row_index, column=len(questions) + 2, value=attempt.raw_score
        ).alignment = CENTER

    _auto_width(sheet, [28] + [10] * len(questions) + [8])
    sheet.freeze_panes = "B2"


def _add_items_sheet(workbook: Workbook, exam: Exam) -> None:
    """Savollar bo'yicha statistika (qiyinlik, to'g'ri javob ulushi)."""
    sheet = workbook.create_sheet("Savollar")

    headers = [
        "№", "Turi", "Bo‘lim", "Kalit", "Qiyinlik b",
        "To‘g‘ri javob ulushi", "Nuqtaviy-biserial",
    ]
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=index, value=header)
    _style_header(sheet, 1, len(headers))

    statistics = getattr(exam, "statistics", None)
    item_map: dict[tuple[int, str], dict] = {}
    if statistics:
        for item in statistics.item_statistics or []:
            item_map[(int(item.get("order", 0)), str(item.get("part", "a")))] = item

    row = 2
    for question in exam.questions.filter(is_active=True).order_by("order"):
        if question.kind == Question.Kind.OPEN:
            key = "; ".join(v for v in (question.answer_a, question.answer_b) if v)
        else:
            key = question.correct_key
        stats_a = item_map.get((question.order, "a"), {})
        values = [
            question.order,
            question.get_kind_display(),
            question.section or "",
            key,
            round(question.difficulty, 4),
            stats_a.get("p_value", ""),
            stats_a.get("point_biserial", ""),
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.border = BORDER
            cell.alignment = LEFT if column in {2, 3, 4} else CENTER
        row += 1

    _auto_width(sheet, [6, 22, 14, 20, 12, 18, 18])
    sheet.freeze_panes = "A2"


# --------------------------------------------------------------------------
#  3. Ishtirokchilar
# --------------------------------------------------------------------------


def participants_workbook(exam: Exam | None = None) -> bytes:
    """Ishtirokchilar ro'yxatini Excel ko'rinishida qaytaradi."""
    from apps.users.models import BotUser

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ishtirokchilar"

    title = f"«{exam.title}» ishtirokchilari" if exam else "Barcha foydalanuvchilar"
    sheet["A1"] = title
    sheet["A1"].font = TITLE_FONT
    sheet.merge_cells("A1:F1")

    headers = ["№", "Ism-familiya", "Telefon", "Telegram ID", "Username", "Ro‘yxatdan o‘tgan"]
    for index, header in enumerate(headers, start=1):
        sheet.cell(row=3, column=index, value=header)
    _style_header(sheet, 3, len(headers))

    if exam is not None:
        users = [
            attempt.user
            for attempt in Attempt.objects.filter(
                exam=exam, status=Attempt.Status.SUBMITTED
            ).select_related("user").ranked()
        ]
    else:
        users = list(BotUser.objects.registered().order_by("-created_at"))

    for row_index, user in enumerate(users, start=1):
        row = row_index + 3
        registered = (
            timezone.localtime(user.registered_at).strftime("%d.%m.%Y %H:%M")
            if user.registered_at
            else ""
        )
        values = [
            row_index,
            user.full_name or user.display_name,
            user.phone or "",
            user.telegram_id,
            f"@{user.username}" if user.username else "",
            registered,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.border = BORDER
            cell.alignment = LEFT if column in {2, 5} else CENTER

    _auto_width(sheet, [6, 28, 16, 14, 18, 20])
    sheet.freeze_panes = "A4"
    return _save(workbook)


def default_filename(prefix: str, exam: Exam | None = None) -> str:
    """Eksport fayli uchun nom yasaydi."""
    from core.text_utils import slugify_filename

    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    if exam is not None:
        return f"{prefix}_{slugify_filename(exam.code)}_{stamp}.xlsx"
    return f"{prefix}_{stamp}.xlsx"


__all__ = [
    "overall_results_workbook",
    "codes_workbook",
    "results_workbook",
    "participants_workbook",
    "default_filename",
]
